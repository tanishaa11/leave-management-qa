import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Bug Report Log"

HEADER_FILL = PatternFill(start_color="7F1D1D", end_color="7F1D1D", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
CRIT_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
HIGH_FILL = PatternFill(start_color="FFE5B4", end_color="FFE5B4", fill_type="solid")
MED_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
BODY_FONT = Font(name="Arial", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLS = ["Bug ID", "Module", "Summary", "Steps to Reproduce", "Expected Result",
        "Actual Result", "Severity", "Priority", "Status", "Linked TC ID", "Date Found"]
WIDTHS = [9, 14, 26, 40, 26, 30, 10, 10, 10, 12, 12]

for c, (h, w) in enumerate(zip(COLS, WIDTHS), start=1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = BORDER
    ws.column_dimensions[cell.column_letter].width = w
ws.freeze_panes = "A2"

bugs = [
["BUG-01", "Employee API", "Duplicate email crashes the server instead of a clean validation error",
 "1) Create an employee with email X\n2) Create another employee with the same email X",
 "API returns 400 with message like 'Email already registered'",
 "API returns 500 Internal Server Error (unhandled DB unique-constraint exception)",
 "Medium", "Medium", "Open", "TC-02", "2026-09-23"],

["BUG-02", "Leave Request API", "Leave request accepted with end date earlier than start date",
 "POST /leave-requests with start_date=2026-10-10, end_date=2026-10-05",
 "API rejects with 400 'end_date must be on/after start_date'",
 "Request is accepted and stored as a valid PENDING request",
 "High", "High", "Open", "TC-04", "2026-09-23"],

["BUG-03", "Approval Workflow", "No status guard on approve - requests can be approved twice, or approved after rejection",
 "1) Approve a PENDING request\n2) Call approve again on the same id\n"
 "-- OR --\n1) Reject a PENDING request\n2) Call approve on the same id",
 "API returns 400 'request already processed' in both cases; balance unaffected",
 "Request is (re-)approved successfully and leave balance is deducted again",
 "Critical", "Critical", "Open", "TC-08, TC-10", "2026-09-23"],

["BUG-04", "Approval Workflow", "Leave balance not validated before approval - can go negative",
 "1) Employee has 10 days remaining\n2) Submit and approve a 30-day leave request",
 "API rejects with 400 'insufficient leave balance'",
 "Request is approved; employee's leave_balance becomes -20",
 "Critical", "Critical", "Open", "TC-09", "2026-09-23"],

["BUG-05", "Leave Request API", "No duplicate/overlap check on leave request submission",
 "1) Submit a leave request for 2026-10-01 to 2026-10-05\n"
 "2) Submit another request for the same employee, same dates",
 "API rejects with 400/409, or flags the overlap for review",
 "Duplicate request is accepted with no warning or conflict check",
 "Medium", "Low", "Open", "TC-11", "2026-09-23"],
]

sev_fill = {"Critical": CRIT_FILL, "High": HIGH_FILL, "Medium": MED_FILL}

for r, row in enumerate(bugs, start=2):
    for c, val in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.font = BODY_FONT
        cell.alignment = WRAP
        cell.border = BORDER
    sev_cell = ws.cell(row=r, column=7)
    if sev_cell.value in sev_fill:
        sev_cell.fill = sev_fill[sev_cell.value]

wb.save("/home/claude/leave-management-system/QA_Documentation/Bug_Report_Log.xlsx")
print("saved")
