import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Bug Fix Report"

HEADER_FILL = PatternFill(start_color="14532D", end_color="14532D", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
CLOSED_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
BODY_FONT = Font(name="Arial", size=10)
CODE_FONT = Font(name="Consolas", size=9)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLS = ["Bug ID", "Summary", "Root Cause", "Fix Applied", "File Changed",
        "Retest (TC ID)", "Retest Result", "Status", "Fixed Date"]
WIDTHS = [9, 24, 30, 38, 12, 12, 26, 10, 12]

for c, (h, w) in enumerate(zip(COLS, WIDTHS), start=1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = BORDER
    ws.column_dimensions[cell.column_letter].width = w
ws.freeze_panes = "A2"

DATE = "2026-09-23"

rows = [
["BUG-01", "Duplicate email crashed the server (500) instead of a clean error",
 "No pre-check for an existing email before insert; the DB's unique-constraint "
 "error was unhandled and bubbled up as a 500.",
 "Added a query for an existing employee by email before insert; returns 400 "
 "'Email already registered' if found.",
 "main.py", "TC-02", "Re-ran duplicate-email request -> 400 with clear message (was 500)", "Closed", DATE],

["BUG-02", "Leave request accepted with end date before start date",
 "No validation comparing start_date and end_date on request creation.",
 "Added a check: if end_date < start_date, return 400 'end_date must be on or "
 "after start_date'.",
 "main.py", "TC-04", "Re-ran inverted-date request -> 400 (was 200, silently accepted)", "Closed", DATE],

["BUG-03", "No status guard on approve/reject - requests could be approved twice, "
 "or approved after rejection, deducting balance again",
 "approve()/reject() updated req.status unconditionally, with no check of the "
 "request's current status.",
 "Added a guard on both approve() and reject(): only a PENDING request can be "
 "approved or rejected; anything else returns 400.",
 "main.py", "TC-08, TC-10", "Re-approved an APPROVED request -> 400. Approved a REJECTED "
 "request -> 400. Balance no longer double-deducted.", "Closed", DATE],

["BUG-04", "Leave balance not validated before approval - could go negative",
 "approve() deducted employee.leave_balance -= days with no check that enough "
 "balance existed.",
 "Added a check before deduction: if requested days > employee.leave_balance, "
 "return 400 'Insufficient leave balance'.",
 "main.py", "TC-09", "Approved a 30-day request against a 20-day balance -> 400 "
 "(previously succeeded, balance went to -10)", "Closed", DATE],

["BUG-05", "Duplicate/overlapping leave requests allowed for the same employee",
 "create_leave_request() had no check against the employee's existing requests "
 "before inserting a new one.",
 "Added an overlap query against the employee's non-REJECTED requests "
 "(start_date <= new.end_date AND end_date >= new.start_date); returns 409 if "
 "an overlap is found.",
 "main.py", "TC-11", "Submitted an overlapping request -> 409 with the conflicting "
 "request's ID and dates (previously accepted silently)", "Closed", DATE],
]

for r, row in enumerate(rows, start=2):
    for c, val in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.font = CODE_FONT if c == 4 else BODY_FONT
        cell.alignment = WRAP
        cell.border = BORDER
    ws.cell(row=r, column=8).fill = CLOSED_FILL

wb.save("/home/claude/leave-management-system/QA_Documentation/Bug_Fix_Report.xlsx")
print("saved")
