import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
PASS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
PENDING_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
BODY_FONT = Font(name="Arial", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLS = ["TC ID", "Module", "Test Scenario", "Test Steps", "Test Data",
        "Expected Result", "Actual Result", "Status", "Linked Bug ID"]
WIDTHS = [8, 14, 26, 34, 22, 30, 30, 10, 12]


def style_sheet(ws, rows):
    for c, (header, w) in enumerate(zip(COLS, WIDTHS), start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A2"
    for r, row in enumerate(rows, start=2):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = BORDER
        status_cell = ws.cell(row=r, column=8)
        if status_cell.value == "Pass":
            status_cell.fill = PASS_FILL
        elif status_cell.value == "Fail":
            status_cell.fill = FAIL_FILL
        else:
            status_cell.fill = PENDING_FILL


# ---------------- Functional ----------------
functional = [
["TC-01", "Employee", "Create employee with valid data",
 "POST /employees with valid name, email, leave_balance",
 "name=Aman Verma, email=aman@test.com, balance=20",
 "201/200 response; employee created with given balance",
 "Employee created successfully with id=1", "Pass", ""],

["TC-02", "Employee", "Create employee with an email already in use",
 "POST /employees using an email that already exists",
 "email=aman@test.com (duplicate)",
 "400 response with a clear 'email already exists' message",
 "500 Internal Server Error (unhandled DB constraint error)", "Fail", "BUG-01"],

["TC-03", "Leave Request", "Submit a valid leave request",
 "POST /leave-requests with a 5-day range",
 "start=2026-10-01, end=2026-10-05",
 "Request created with status PENDING",
 "Request created, status PENDING", "Pass", ""],

["TC-04", "Leave Request", "Submit request with end date before start date",
 "POST /leave-requests with an inverted date range",
 "start=2026-10-10, end=2026-10-05",
 "400 validation error; request rejected",
 "Request accepted and stored as PENDING", "Fail", "BUG-02"],

["TC-05", "Leave Request", "Submit leave request for a non-existent employee",
 "POST /leave-requests with an invalid employee_id",
 "employee_id=999",
 "404 'Employee not found'",
 "404 'Employee not found'", "Pass", ""],

["TC-06", "Employee", "Fetch leave balance before any approval",
 "GET /employees/{id}",
 "employee_id=1",
 "Returns unchanged starting balance (20)",
 "Returned balance=20", "Pass", ""],

["TC-07", "Approval", "Approve a pending leave request",
 "PUT /leave-requests/{id}/approve",
 "request id=1 (5-day request)",
 "Status -> APPROVED; balance reduced by 5 (20->15)",
 "Status APPROVED; balance 20->15", "Pass", ""],

["TC-08", "Approval", "Re-approve an already-approved leave request",
 "Call approve again on the same request id",
 "request id=1 (already APPROVED)",
 "400 error; request already processed, no further deduction",
 "Approved again; balance deducted a 2nd time (15->10)", "Fail", "BUG-03"],

["TC-09", "Approval", "Approve a request that exceeds remaining balance",
 "Approve a 30-day request when ~10 days remain",
 "request covering 2026-11-01 to 2026-11-30",
 "400 error; 'insufficient leave balance'",
 "Approved anyway; balance went negative (10->-20)", "Fail", "BUG-04"],

["TC-10", "Approval", "Approve a request that was already REJECTED",
 "Reject a request, then call approve on the same id",
 "request id=1",
 "400 error; cannot approve a rejected request",
 "Request approved successfully; balance deducted", "Fail", "BUG-03"],

["TC-11", "Leave Request", "Submit a duplicate/overlapping request",
 "POST a request with identical dates to an existing one, same employee",
 "start=2026-10-01, end=2026-10-05 (duplicate of TC-03)",
 "400/409 error, or flagged as overlapping",
 "Duplicate request accepted with no warning", "Fail", "BUG-05"],

["TC-12", "Rejection", "Reject a pending leave request",
 "PUT /leave-requests/{id}/reject",
 "request id=1 (PENDING)",
 "Status -> REJECTED; leave balance unaffected",
 "Status REJECTED; balance unchanged", "Pass", ""],

["TC-13", "Employee", "Fetch a non-existent employee",
 "GET /employees/999", "employee_id=999",
 "404 'Employee not found'",
 "404 'Employee not found'", "Pass", ""],

["TC-14", "Leave Request", "Fetch a non-existent leave request",
 "GET /leave-requests/999", "request_id=999",
 "404 'Leave request not found'",
 "404 'Leave request not found'", "Pass", ""],
]

ws1 = wb.active
ws1.title = "Functional"
style_sheet(ws1, functional)

# ---------------- Regression (planned re-run after fixes) ----------------
regression = [
["TC-02", "Employee", "Re-test: duplicate email handling after fix",
 "POST /employees with a duplicate email",
 "email=aman@test.com",
 "400 with clear validation message",
 "Not executed - pending BUG-01 fix", "Blocked", "BUG-01"],
["TC-04", "Leave Request", "Re-test: inverted date range after fix",
 "POST /leave-requests with end date before start date",
 "start=2026-10-10, end=2026-10-05",
 "400 validation error",
 "Not executed - pending BUG-02 fix", "Blocked", "BUG-02"],
["TC-08/10", "Approval", "Re-test: status guard on approve after fix",
 "Approve an already-approved and an already-rejected request",
 "request id in APPROVED / REJECTED state",
 "400 error in both cases; no balance change",
 "Not executed - pending BUG-03 fix", "Blocked", "BUG-03"],
["TC-09", "Approval", "Re-test: balance validation on approve after fix",
 "Approve a request exceeding remaining balance",
 "30-day request, ~10 days remaining",
 "400 'insufficient leave balance'",
 "Not executed - pending BUG-04 fix", "Blocked", "BUG-04"],
["TC-11", "Leave Request", "Re-test: overlap/duplicate detection after fix",
 "Submit a request duplicating an existing one",
 "same employee, same date range",
 "400/409 duplicate/overlap error",
 "Not executed - pending BUG-05 fix", "Blocked", "BUG-05"],
["TC-01", "Employee", "Confirm valid employee creation still works",
 "POST /employees with valid data", "name/email/balance",
 "Employee created as before",
 "Pass (unaffected by fixes)", "Pass", ""],
["TC-07", "Approval", "Confirm normal single-approve flow still works",
 "Approve one pending request once", "request id=1",
 "Status APPROVED, balance deducted once",
 "Pass (unaffected by fixes)", "Pass", ""],
]
ws2 = wb.create_sheet("Regression")
style_sheet(ws2, regression)

# ---------------- Smoke ----------------
smoke = [
["SM-01", "Build", "API server starts successfully",
 "Run uvicorn main:app and check /docs loads",
 "n/a", "Server responds 200 on /docs",
 "Server up, /docs returned 200", "Pass", ""],
["SM-02", "Employee", "Create an employee (core path)",
 "POST /employees with minimal valid payload",
 "name, email", "Employee created, id returned",
 "Employee created, id=1", "Pass", ""],
["SM-03", "Leave Request", "Submit a leave request (core path)",
 "POST /leave-requests with valid dates",
 "employee_id=1, valid date range",
 "Request created, status PENDING",
 "Request created, status PENDING", "Pass", ""],
["SM-04", "Approval", "Approve a leave request (core path)",
 "PUT /leave-requests/{id}/approve",
 "a fresh PENDING request",
 "Status APPROVED, balance updated",
 "Status APPROVED, balance updated", "Pass", ""],
["SM-05", "Employee", "Fetch employee record (core path)",
 "GET /employees/{id}", "employee_id=1",
 "Employee record returned",
 "Employee record returned", "Pass", ""],
]
ws3 = wb.create_sheet("Smoke")
style_sheet(ws3, smoke)

wb.save("/home/claude/leave-management-system/QA_Documentation/Test_Cases.xlsx")
print("saved")
